"""Dosya listeleme, belge okuma ve düz metin yazma."""

from __future__ import annotations

import json

from .. import safety
from . import ToolError, ToolResult, rel, resolve_path, safe_path, tool

MAX_CHARS = 60_000


def _sensitive(path: str) -> str | None:
    from .system import sensitive_path  # system modülü files'tan sonra yüklenir

    return sensitive_path(path)


def _truncate(text: str) -> str:
    if len(text) <= MAX_CHARS:
        return text
    return (
        text[:MAX_CHARS]
        + f"\n\n[... İçerik {len(text):,} karakter; ilk {MAX_CHARS:,} karakter gösterildi. "
        "Kalanı için run_python ile parça parça okuyun.]"
    )


@tool(
    "list_files",
    "Çalışma alanındaki dosyaları listeler. Kullanıcının yüklediği dosyalar 'uploads/', "
    "ADA'nın ürettikleri 'outputs/' altındadır.",
    {"directory": {"type": "string", "description": "Göreli klasör yolu; boş bırakılırsa kök."}},
)
def list_files(directory: str = "") -> str:
    base = safe_path(directory, must_exist=True)
    if not base.is_dir():
        raise ToolError(f"Klasör değil: {directory}")
    lines = []
    for p in sorted(base.rglob("*")):
        if p.is_file() and ".ada" not in p.parts:
            lines.append(f"{rel(p)}  ({p.stat().st_size:,} bayt)")
    return "\n".join(lines[:500]) or "(boş)"


def _read_pdf(path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [f"--- Sayfa {i + 1} ---\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages)]
    return "\n".join(pages)


def _read_docx(path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for ti, table in enumerate(doc.tables):
        parts.append(f"\n[Tablo {ti + 1}]")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _read_pptx(path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides):
        out.append(f"--- Slayt {i + 1} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)


def _read_table(path) -> str:
    import pandas as pd

    if path.suffix.lower() == ".csv":
        sheets = {"csv": pd.read_csv(path)}
    else:
        sheets = pd.read_excel(path, sheet_name=None)
    out = []
    for name, df in sheets.items():
        out.append(
            f"=== Sayfa: {name} | {df.shape[0]} satır x {df.shape[1]} sütun ===\n"
            f"Sütunlar ve tipler:\n{df.dtypes.to_string()}\n\n"
            f"İlk 30 satır:\n{df.head(30).to_string()}"
        )
        numeric = df.select_dtypes("number")
        if not numeric.empty:
            out.append(f"\nSayısal özet:\n{numeric.describe().to_string()}")
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        out.append(_formula_summary(path))
    out.append("\n(Detaylı analiz için run_python aracında pandas kullanın.)")
    return "\n".join(out)


def _formula_summary(path) -> str:
    """pandas formül hücrelerini NaN gösterir (Excel açılana kadar hesaplanmış değer yoktur);
    formülleri ayrıca listeleyerek dosyanın doğru kurulduğunun kontrol edilebilmesini sağlar."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=False, read_only=True)
    except Exception as e:  # noqa: BLE001
        return f"\n(Formüller okunamadı: {e})"
    lines, total = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    total += 1
                    if len(lines) < 40:
                        lines.append(f"  {ws.title}!{cell.coordinate}: {cell.value}")
    wb.close()
    if not total:
        return "\nFormül yok (tüm hücreler sabit değer)."
    more = f"\n  … toplam {total} formül" if total > len(lines) else ""
    return (
        f"\nFormüller ({total} hücre; yukarıdaki tabloda NaN görünen hücreler bunlardır, "
        "değerleri Excel'de açılınca hesaplanır):\n" + "\n".join(lines) + more
    )


@tool(
    "read_document",
    "Bir belgenin içeriğini metin olarak okur. PDF, Word (.docx), PowerPoint (.pptx), "
    "Excel (.xlsx/.xls), CSV, TXT, MD, JSON ve benzeri metin dosyalarını destekler. "
    "Tablolarda şema, ilk satırlar ve sayısal özet döner. Bilgisayardaki her yol geçerlidir.",
    {"path": {"type": "string", "description": "uploads/rapor.pdf gibi göreli ya da ~/Desktop/x.xlsx gibi mutlak yol"}},
    ["path"],
    risk=lambda a: _sensitive(a.get("path", "")),
)
def read_document(path: str) -> str:
    p = resolve_path(path, must_exist=True)
    if p.is_dir():
        raise ToolError(f"Bu bir klasör: {p}")
    suffix = p.suffix.lower()
    readers = {
        ".pdf": _read_pdf,
        ".docx": _read_docx,
        ".pptx": _read_pptx,
        ".xlsx": _read_table,
        ".xlsm": _read_table,
        ".xls": _read_table,
        ".csv": _read_table,
    }
    if suffix in readers:
        return _truncate(readers[suffix](p))
    try:
        return _truncate(p.read_text(encoding="utf-8"))
    except UnicodeDecodeError:
        raise ToolError(f"'{suffix}' uzantılı ikili dosya okunamıyor. run_python ile deneyin.")


@tool(
    "write_text_file",
    "outputs/ klasörüne düz metin dosyası yazar (.md, .txt, .csv, .json, .html vb.).",
    {
        "filename": {"type": "string", "description": "Uzantılı dosya adı, ör. ozet.md"},
        "content": {"type": "string"},
    },
    ["filename", "content"],
)
def write_text_file(filename: str, content: str) -> ToolResult:
    p = safe_path(f"outputs/{filename}")
    p.parent.mkdir(parents=True, exist_ok=True)
    safety.backup(p)
    if p.suffix.lower() == ".json":
        try:
            content = json.dumps(json.loads(content), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass
    p.write_text(content, encoding="utf-8")
    return ToolResult(f"Kaydedildi: {rel(p)}", files=[rel(p)])
