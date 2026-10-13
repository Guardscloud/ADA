"""Word, Excel, PowerPoint ve grafik üretimi."""

from __future__ import annotations

import re

from . import ToolError, ToolResult, output_path, rel, tool

_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _add_runs(paragraph, text: str) -> None:
    """**kalın** ve *italik* işaretlerini Word biçimine çevirir."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _add_table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        for j in range(width):
            cell = table.cell(i, j)
            cell.text = ""
            _add_runs(cell.paragraphs[0], row[j] if j < len(row) else "")
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


@tool(
    "create_word_document",
    "Markdown içerikten biçimli bir Word (.docx) belgesi oluşturur: rapor, mektup, "
    "sözleşme taslağı, toplantı tutanağı, teklif vb. Desteklenen markdown: # başlıklar, "
    "- madde işaretleri, 1. numaralı liste, **kalın**, *italik*, | tablolar |, --- sayfa sonu.",
    {
        "filename": {"type": "string", "description": "Dosya adı, ör. aylik_rapor.docx"},
        "title": {"type": "string", "description": "Belge başlığı (kapakta gösterilir)"},
        "markdown": {"type": "string", "description": "Belgenin gövdesi (markdown)"},
    },
    ["filename", "markdown"],
)
def create_word_document(filename: str, markdown: str, title: str = "") -> ToolResult:
    import docx
    from docx.enum.text import WD_BREAK

    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue
        if not stripped:
            pass
        elif stripped == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif m := re.match(r"^(#{1,6})\s+(.*)", stripped):
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        elif m := re.match(r"^[-*•]\s+(.*)", stripped):
            _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\d+[.)]\s+(.*)", stripped):
            _add_runs(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            _add_runs(doc.add_paragraph(), stripped)
        i += 1

    path = output_path(filename, ".docx")
    doc.save(path)
    return ToolResult(f"Word belgesi oluşturuldu: {rel(path)}", files=[rel(path)])


@tool(
    "create_excel",
    "Bir veya birden çok sayfalı Excel (.xlsx) dosyası oluşturur. Başlık satırı kalın, "
    "sütun genişlikleri otomatik. '=' ile başlayan hücreler Excel formülü olarak yazılır "
    "(ör. '=SUM(B2:B10)').",
    {
        "filename": {"type": "string"},
        "sheets": {
            "type": "array",
            "description": "Sayfalar",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "rows": {
                        "type": "array",
                        "description": "İlk satır başlık olmak üzere satır listesi",
                        "items": {"type": "array", "items": {}},
                    },
                },
                "required": ["name", "rows"],
            },
        },
    },
    ["filename", "sheets"],
)
def create_excel(filename: str, sheets: list[dict]) -> ToolResult:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not sheets:
        raise ToolError("En az bir sayfa gerekli.")
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(str(sheet.get("name") or "Sayfa")[:31])
        for row in sheet.get("rows", []):
            ws.append(list(row))
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F5597")
            ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 60)
    path = output_path(filename, ".xlsx")
    wb.save(path)
    report = _excel_review(wb)
    return ToolResult(f"Excel dosyası oluşturuldu: {rel(path)}\n\n{report}", files=[rel(path)])


_SUM_RANGE = re.compile(r"^=\s*(SUM|TOPLA|AVERAGE|ORTALAMA|MIN|MAX)\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*$", re.I)


def _excel_review(wb) -> str:
    """Kaydedilen tablonun önizlemesi + tipik formül kayması uyarıları (model kendi işini düzeltebilsin)."""
    from openpyxl.utils import get_column_letter

    parts, warnings = [], []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15)):
            rows.append(" | ".join("" if c.value is None else str(c.value)[:18] for c in r[:10]))
        parts.append(f"Önizleme [{ws.title}] ({ws.max_row} satır x {ws.max_column} sütun):\n" + "\n".join(rows))
        for row in ws.iter_rows():
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                m = _SUM_RANGE.match(cell.value)
                if not m:
                    continue
                c1, r1, c2, r2 = m.group(2).upper(), int(m.group(3)), m.group(4).upper(), int(m.group(5))
                col = get_column_letter(cell.column)
                if c1 == c2 and c1 != col:
                    warnings.append(f"{ws.title}!{cell.coordinate} {cell.value}: kendi sütunu ({col}) yerine {c1} sütununu topluyor — kayma olabilir.")
                elif r1 == r2 and c1 != c2 and r1 != cell.row:  # yatay aralık başka satırda
                    warnings.append(f"{ws.title}!{cell.coordinate} {cell.value}: kendi satırı ({cell.row}) yerine {r1}. satırı topluyor — kayma olabilir.")
                elif c1 == c2 and r1 <= cell.row <= r2:
                    warnings.append(f"{ws.title}!{cell.coordinate} {cell.value}: kendi hücresini kapsıyor (döngüsel başvuru).")
    out = "\n\n".join(parts)
    if warnings:
        out += "\n\n⚠ FORMÜL KONTROLÜ — düzeltip dosyayı yeniden oluştur:\n" + "\n".join(f"- {w}" for w in warnings[:20])
    else:
        out += "\n\nFormül kontrolü: kayma veya döngüsel başvuru bulunmadı."
    return out


@tool(
    "create_presentation",
    "PowerPoint (.pptx) sunumu oluşturur. Her slayt başlık, madde işaretleri, isteğe bağlı "
    "konuşmacı notu ve isteğe bağlı görsel (ör. create_chart ile üretilmiş outputs/grafik.png) içerir.",
    {
        "filename": {"type": "string"},
        "title": {"type": "string", "description": "Kapak slaydı başlığı"},
        "subtitle": {"type": "string"},
        "slides": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "bullets": {"type": "array", "items": {"type": "string"}},
                    "notes": {"type": "string"},
                    "image": {"type": "string", "description": "Çalışma alanına göre görsel yolu"},
                },
                "required": ["title"],
            },
        },
    },
    ["filename", "title", "slides"],
)
def create_presentation(filename: str, title: str, slides: list[dict], subtitle: str = "") -> ToolResult:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    from . import safe_path

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title
    if subtitle and len(cover.placeholders) > 1:
        cover.placeholders[1].text = subtitle

    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s.get("title", "")
        body = slide.placeholders[1]
        image = s.get("image")
        if image:
            body.width = Inches(6.2)
        tf = body.text_frame
        tf.clear()
        for j, bullet in enumerate(s.get("bullets") or []):
            para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            para.text = bullet
            para.font.size = Pt(20)
        if image:
            slide.shapes.add_picture(str(safe_path(image, must_exist=True)), Inches(7), Inches(1.6), width=Inches(5.8))
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    path = output_path(filename, ".pptx")
    prs.save(path)
    return ToolResult(f"Sunum oluşturuldu ({len(slides) + 1} slayt): {rel(path)}", files=[rel(path)])


@tool(
    "create_chart",
    "Veriden PNG grafik üretir (çizgi, sütun, yatay sütun, pasta). Raporlara ve sunumlara "
    "eklenebilir. Daha karmaşık grafikler için run_python + matplotlib kullanın.",
    {
        "filename": {"type": "string"},
        "chart_type": {"type": "string", "enum": ["line", "bar", "barh", "pie"]},
        "title": {"type": "string"},
        "labels": {"type": "array", "items": {"type": "string"}, "description": "X ekseni etiketleri"},
        "series": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "values": {"type": "array", "items": {"type": "number"}},
                },
                "required": ["name", "values"],
            },
        },
        "x_label": {"type": "string"},
        "y_label": {"type": "string"},
    },
    ["filename", "chart_type", "title", "labels", "series"],
)
def create_chart(
    filename: str,
    chart_type: str,
    title: str,
    labels: list[str],
    series: list[dict],
    x_label: str = "",
    y_label: str = "",
) -> ToolResult:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=150)
    if chart_type == "pie":
        ax.pie(series[0]["values"], labels=labels, autopct="%1.1f%%", startangle=90)
        ax.axis("equal")
    else:
        x = np.arange(len(labels))
        width = 0.8 / max(len(series), 1)
        for k, s in enumerate(series):
            if chart_type == "line":
                ax.plot(labels, s["values"], marker="o", linewidth=2, label=s["name"])
            elif chart_type == "bar":
                ax.bar(x + k * width - 0.4 + width / 2, s["values"], width, label=s["name"])
            else:
                ax.barh(x + k * width - 0.4 + width / 2, s["values"], width, label=s["name"])
        if chart_type == "bar":
            ax.set_xticks(x, labels, rotation=30 if len(labels) > 6 else 0, ha="right" if len(labels) > 6 else "center")
        elif chart_type == "barh":
            ax.set_yticks(x, labels)
        ax.set_xlabel(x_label)
        ax.set_ylabel(y_label)
        ax.grid(axis="y" if chart_type != "barh" else "x", alpha=0.3)
        ax.spines[["top", "right"]].set_visible(False)
        if len(series) > 1:
            ax.legend(frameon=False)
    ax.set_title(title, fontsize=14, fontweight="bold")
    fig.tight_layout()
    path = output_path(filename, ".png")
    fig.savefig(path)
    plt.close(fig)
    return ToolResult(f"Grafik oluşturuldu: {rel(path)}", files=[rel(path)])
